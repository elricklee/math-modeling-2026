"""附件精细读数：完整精度导出 + 数值体检。

解决的问题：此前的结构体检用 `%g` 格式化浮点，只保留 6 位有效数字，
导致 3439.8466 被显示成 3439.85。本脚本一律使用 `repr`（Python 3 最短往返表示），
保证不丢任何精度。

产出：
    results/attachment_precision.md   —— 完整精度报告（含附件 1 全部 144 行）

运行：
    python -m src.inspect_precision
"""

from __future__ import annotations

import sys
from datetime import datetime, time
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

from src.config import RAW_DIR, RESULTS_DIR

ATTACH_DIR = RAW_DIR / "C题"
TEMPLATE_DIR = ATTACH_DIR / "附件5"


def fmt(v: object) -> str:
    """完整精度格式化。"""
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d %H:%M")
    if isinstance(v, time):
        return v.strftime("%H:%M")
    if isinstance(v, float):
        return repr(v)
    return str(v)


def decimals(v: object) -> int:
    """该数值的小数位数（用于统计数据精细度）。"""
    if not isinstance(v, (int, float)):
        return 0
    d = Decimal(repr(float(v))).normalize()
    e = -d.as_tuple().exponent
    return max(e, 0)


def grid(path: Path, sheet: str) -> list[list[object]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        return [list(r) for r in wb[sheet].iter_rows(values_only=True)]
    finally:
        wb.close()


def sheets(path: Path) -> list[str]:
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def col_profile(rows: list[list[object]], col: int, skip_header: int = 1) -> str:
    """对某列做数值体检：个数、小数位、范围、负值、缺失。"""
    vals = [r[col] for r in rows[skip_header:] if col < len(r)]
    nums = [v for v in vals if isinstance(v, (int, float))]
    miss = sum(1 for v in vals if v is None)
    neg = sum(1 for v in nums if v < 0)
    dec = max((decimals(v) for v in nums), default=0)
    if not nums:
        return f"非数值列（{len(vals)} 个单元格，缺失 {miss}）"
    return (f"n={len(nums)} 缺失={miss} 最大小数位={dec} "
            f"min={repr(min(nums))} max={repr(max(nums))} 负值={neg}")


def main() -> None:
    L: list[str] = []
    add = L.append

    add("# C 题附件精细读数报告（完整精度）")
    add("")
    add("> 由 `python -m src.inspect_precision` 生成。")
    add("> 所有数值均以 `repr(float)` 输出（Python 3 最短往返表示），**不截断、不四舍五入**。")
    add("")

    # ============ 附件 1：逐行完整数值 ============
    a1 = grid(ATTACH_DIR / "附件1.xlsx", "Sheet1")
    add("## 一、附件 1 全部数据（计算粒度 10 分钟，共 %d 个区间）" % (len(a1) - 1))
    add("")
    add("表头：" + " | ".join(fmt(v) for v in a1[0]))
    add("")
    add("数值体检：")
    for c, name in ((1, "电价(元/kWh)"), (2, "小区负载(kW)"), (3, "光伏预测功率(kW)")):
        add(f"- 第 {c} 列 {name}: {col_profile(a1, c)}")
    add("")
    add("| # | 时间 | 电价 | 小区负载 | 光伏预测功率 |")
    add("|---:|---|---:|---:|---:|")
    for i, r in enumerate(a1[1:], start=1):
        add(f"| {i} | {fmt(r[0])} | {fmt(r[1])} | {fmt(r[2])} | {fmt(r[3])} |")
    add("")

    # ============ 附件 2 ============
    p2 = ATTACH_DIR / "附件2.xlsx"
    add("## 二、附件 2（负载 + 光伏实际）")
    add("")
    for sh in sheets(p2):
        rows = grid(p2, sh)
        add(f"### {sh}  形状 = {len(rows)} 行 × {len(rows[0])} 列")
        add("")
        add("行标签（列 A）起止：" + f"{fmt(rows[1][0])} → {fmt(rows[-1][0])}"
            + f"（共 {len(rows) - 1} 天）")
        add("")
        add("列标签（第 1 行）前 4 个：" + " | ".join(fmt(v) for v in rows[0][:4]))
        add("末 2 个：" + " | ".join(fmt(v) for v in rows[0][-2:]))
        add("")
        prof = [col_profile(rows, c) for c in range(1, len(rows[0]))]
        add(f"- 全部 {len(prof)} 个数值列的最大小数位 = "
            f"{max(int(s.split('最大小数位=')[1].split(' ')[0]) for s in prof if '最大小数位=' in s)}")
        mins = [min(r[c] for r in rows[1:] if isinstance(r[c], (int, float)))
                for c in range(1, len(rows[0]))]
        maxs = [max(r[c] for r in rows[1:] if isinstance(r[c], (int, float)))
                for c in range(1, len(rows[0]))]
        neg = sum(1 for r in rows[1:] for c in range(1, len(rows[0]))
                  if isinstance(r[c], (int, float)) and r[c] < 0)
        miss = sum(1 for r in rows[1:] for c in range(1, len(rows[0])) if r[c] is None)
        add(f"- 全表 min={repr(min(mins))}  max={repr(max(maxs))}  "
            f"负值={neg}  缺失={miss}")
        add("")
        add(f"首行（{fmt(rows[1][0])}）完整数值：")
        add("")
        add("```")
        add("  ".join(fmt(v) for v in rows[1][1:]))
        add("```")
        add("")

    # ============ 附件 3 ============
    p3 = ATTACH_DIR / "附件3.xlsx"
    add("## 三、附件 3（光伏预报）")
    add("")
    for sh in sheets(p3):
        rows = grid(p3, sh)
        add(f"### {sh}  形状 = {len(rows)} 行 × {len(rows[0])} 列")
        add("")
        add("完整表头：" + " | ".join(fmt(v) for v in rows[0]))
        add("")
        dates = [r[0] for r in rows[1:] if r[0] is not None]
        pubs = [r[1] for r in rows[1:] if r[1] is not None]
        add(f"- 日期列非空个数 = {len(dates)}，起止 {fmt(dates[0])} → {fmt(dates[-1])}")
        add(f"- 预报时刻非空个数 = {len(pubs)}，取值集合 = {sorted(set(fmt(v) for v in pubs))}")
        add(f"- 数据行数 = {len(rows) - 1}（应为 365 天 × 4 次 = 1460）")
        add("")
        prof = [col_profile(rows, c) for c in range(2, len(rows[0]))]
        decs = [int(s.split("最大小数位=")[1].split(" ")[0]) for s in prof if "最大小数位=" in s]
        neg = sum(1 for r in rows[1:] for c in range(2, len(rows[0]))
                  if isinstance(r[c], (int, float)) and r[c] < 0)
        miss = sum(1 for r in rows[1:] for c in range(2, len(rows[0])) if r[c] is None)
        add(f"- 24 个预报列的最大小数位 = {max(decs) if decs else 0}")
        add(f"- 负值 = {neg}   缺失 = {miss}")
        add("")
        add("前 5 个数据行（完整精度）：")
        add("")
        add("```")
        for r in rows[1:6]:
            add("  " + " | ".join(fmt(v) for v in r))
        add("```")
        add("")

    # ============ 附件 4 ============
    p4 = ATTACH_DIR / "附件4.xlsx"
    add("## 四、附件 4（电价）")
    add("")
    for sh in sheets(p4):
        rows = grid(p4, sh)
        add(f"### {sh}  形状 = {len(rows)} 行 × {len(rows[0])} 列")
        add("")
        add("列标签（第 1 行）前 4 个：" + " | ".join(fmt(v) for v in rows[0][:4]))
        add("末 2 个：" + " | ".join(fmt(v) for v in rows[0][-2:]))
        add("")
        prof = [col_profile(rows, c) for c in range(1, len(rows[0]))]
        decs = [int(s.split("最大小数位=")[1].split(" ")[0]) for s in prof if "最大小数位=" in s]
        minv = min(min(r[c] for r in rows[1:] if isinstance(r[c], (int, float)))
                   for c in range(1, len(rows[0])))
        maxv = max(max(r[c] for r in rows[1:] if isinstance(r[c], (int, float)))
                   for c in range(1, len(rows[0])))
        neg = sum(1 for r in rows[1:] for c in range(1, len(rows[0]))
                  if isinstance(r[c], (int, float)) and r[c] < 0)
        miss = sum(1 for r in rows[1:] for c in range(1, len(rows[0])) if r[c] is None)
        add(f"- 全部 {len(prof)} 个数值列的最大小数位 = {max(decs) if decs else 0}")
        add(f"- min={repr(minv)}  max={repr(maxv)}  负值={neg}  缺失={miss}")
        add("")
        add(f"首行（{fmt(rows[1][0])}）完整数值：")
        add("")
        add("```")
        add("  ".join(fmt(v) for v in rows[1][1:]))
        add("```")
        add("")

    # ============ 附件 5 ============
    add("## 五、附件 5 结果模板")
    add("")
    for f in sorted(TEMPLATE_DIR.glob("result*.xlsx")):
        add(f"### {f.name}")
        add("")
        for sh in sheets(f):
            rows = grid(f, sh)
            add(f"- `{sh}`  {len(rows)} 行 × {len(rows[0])} 列；表头 = "
                + " | ".join(fmt(v) for v in rows[0]))
        add("")

    text = "\n".join(L)
    RESULTS_DIR.mkdir(parents=True, exist_ok=True)
    dest = RESULTS_DIR / "attachment_precision.md"
    dest.write_text(text, encoding="utf-8")
    print(f"[written] {dest}", file=sys.stderr)


if __name__ == "__main__":
    main()
