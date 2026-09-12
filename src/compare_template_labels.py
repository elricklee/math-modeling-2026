"""对齐比对附件 1 的时间轴与附件 5 模板的时间标签。

背景：附件 1 的时间列是单个时刻（00:10 … 24:00），共 144 个；
      result1.xlsx「计划购电量」列 A 是区间标签（0:10-0:20 … ），共 144 个；
      result2.xlsx「计划购电量」表头是区间标签 + 2 列合计，共 144 + 2 列。
三者是否一一对齐，直接决定 result1.xlsx 该往哪一行写数。
"""

from __future__ import annotations

import sys

from openpyxl import load_workbook

from src.config import RAW_DIR, RESULTS_DIR

A1 = RAW_DIR / "C题" / "附件1.xlsx"
T = RAW_DIR / "C题" / "附件5"


def cells(path, sheet, min_row, max_row, min_col, max_col, by_row: bool):
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[sheet]
        if by_row:
            r = next(ws.iter_rows(min_row=min_row, max_row=max_row,
                                  min_col=min_col, max_col=max_col, values_only=True))
            return list(r)
        return [r[0] for r in ws.iter_rows(min_row=min_row, max_row=max_row,
                                           min_col=min_col, max_col=max_col,
                                           values_only=True)]
    finally:
        wb.close()


def s(v) -> str:
    return "" if v is None else str(v)


def main() -> None:
    a1_time = cells(A1, "Sheet1", 2, 145, 1, 1, by_row=False)          # 144 个时刻
    r1_lab = cells(T / "result1.xlsx", "计划购电量", 2, 145, 1, 1, by_row=False)
    r2_hdr = cells(T / "result2.xlsx", "计划购电量", 1, 1, 1, 147, by_row=True)
    r2_time = r2_hdr[1:145]                                            # 144 个区间
    r2_extra = r2_hdr[145:]                                            # 尾部 2 列

    L: list[str] = []
    add = L.append
    add("# 附件 1 时间轴 vs 附件 5 模板时间标签 对齐比对")
    add("")
    add(f"- 附件 1 时间列个数 = {len(a1_time)}，首 = `{s(a1_time[0])}`，末 = `{s(a1_time[-1])}`")
    add(f"- result1 列A 标签数 = {len(r1_lab)}，首 = `{s(r1_lab[0])}`，末 = `{s(r1_lab[-1])}`")
    add(f"- result2 时间列数 = {len(r2_time)}，首 = `{s(r2_time[0])}`，末 = `{s(r2_time[-1])}`")
    add(f"- result2 尾部额外列 = {[s(v) for v in r2_extra]}")
    add("")
    add("## 前 6 项")
    add("")
    add("| idx | 附件1 时刻 | result1 区间标签 | result2 区间标签 |")
    add("|---:|---|---|---|")
    for i in range(6):
        add(f"| {i+1} | {s(a1_time[i])} | {s(r1_lab[i])} | {s(r2_time[i])} |")
    add("")
    add("## 后 8 项")
    add("")
    add("| idx | 附件1 时刻 | result1 区间标签 | result2 区间标签 |")
    add("|---:|---|---|---|")
    n = len(a1_time)
    for i in range(n - 8, n):
        add(f"| {i+1} | {s(a1_time[i])} | {s(r1_lab[i])} | {s(r2_time[i])} |")
    add("")
    add("## result1 与 result2 标签是否逐项相同")
    add("")
    diff = [(i + 1, s(r1_lab[i]), s(r2_time[i])) for i in range(n) if s(r1_lab[i]) != s(r2_time[i])]
    if not diff:
        add("完全相同，共 144 项。")
    else:
        add(f"共 {len(diff)} 项不同：")
        add("")
        add("| idx | result1 | result2 |")
        add("|---:|---|---|")
        for i, x, y in diff:
            add(f"| {i} | {x} | {y} |")
    add("")
    add("## 标签自身是否连续（10 分钟一档）")
    add("")
    add("result1 第 2 项起的标签区间左端点：" + " ".join(s(v).split("-")[0] for v in r1_lab[1:5]))
    add("result1 末 3 项区间：" + " / ".join(s(v) for v in r1_lab[-3:]))
    add("")

    text = "\n".join(L)
    RESULTS_DIR.mkdir(parents=True, exist_ok=True)
    dest = RESULTS_DIR / "label_alignment.md"
    dest.write_text(text, encoding="utf-8")
    sys.stdout.write(text + "\n")


if __name__ == "__main__":
    main()
