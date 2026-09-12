"""Behavioral regression checks for the review's concrete failure cases."""
from datetime import date, timedelta
import tempfile
from pathlib import Path
import unittest
import numpy as np
import openpyxl
from src.CUMCM2026_C.dispatch import plan_segment, execute_interval, settlement, history_prediction, forecast_vector
from src.CUMCM2026_C.export_results import export_scenario


class DispatchTests(unittest.TestCase):
    def test_later_forecast_is_not_used_by_earlier_release(self):
        from types import SimpleNamespace
        day=date(2025,2,1)
        attachment=SimpleNamespace(release_days=[day,day],release_hours=[0,6],
                                   forecast=np.array([[100.]*24,[200.]*24]))
        first=forecast_vector(attachment,day,0)
        attachment.forecast[1]=1e9
        np.testing.assert_array_equal(first,forecast_vector(attachment,day,0))
        self.assertTrue(np.isnan(forecast_vector(attachment,day,6)[:36]).all())

    def test_adjustment_fee_changes_decision(self):
        # An already paid plan may be unused for free; reducing it incurs a fee.
        p=plan_segment(np.ones(2),np.zeros(2),np.zeros(2),6000.,np.full(2,100.))
        np.testing.assert_allclose(p['purchase'],100.,atol=1e-5)
        self.assertLess(p['objective_yuan'],1e-6)

    def test_shortfall_is_emergency_not_spare(self):
        e=execute_interval(6000,0,0,0,600,0)
        self.assertAlmostEqual(e['emergency'],100.)
        self.assertAlmostEqual(e['unused'],0.)
        s=settlement([1],[0],[0],[e['emergency']])
        self.assertAlmostEqual(s['emergency_cost_yuan'],500.)

    def test_supply_execution_is_physical(self):
        state=6000.
        for load,pv in [(0,10000),(10000,0),(100,10000),(0,0)]:
            r=execute_interval(state,100,833.333333,0,load,pv)
            self.assertGreaterEqual(r['soc'],1200-1e-7)
            self.assertLessEqual(r['soc'],10800+1e-7)
            self.assertAlmostEqual(100-r['unused']+r['emergency']+r['discharge']+pv/6-r['curtail'],load/6+r['charge'],places=7)
            state=r['soc']

    def test_future_rows_do_not_change_history_forecast(self):
        dates=[date(2025,1,1)+timedelta(days=i) for i in range(50)]
        a=np.arange(50*144,dtype=float).reshape(50,144)
        before=history_prediction(a,30,dates,.8)
        a[30:]=1e10
        np.testing.assert_array_equal(before,history_prediction(a,30,dates,.8))

    def test_spill_and_minimum_throughput_remove_cycles(self):
        r=plan_segment(np.ones(24),np.full(24,1000.),np.full(24,8000.),6000.)
        self.assertFalse(np.any((r['charge']>1e-6)&(r['discharge']>1e-6)))
        self.assertLess(np.max(r['charge']+r['discharge']),833.333334)

    def test_export_does_not_rotate_and_fills_every_storage_block(self):
        records=[]
        for day in range(3):
            records.append({'date':f'2025-02-{day+1:02d}', 'price_by_interval':[1.]*144,
                            'plan_by_interval_kwh':list(range(144)), 'purchase_by_interval_kwh':list(range(144)),
                            'charge_by_interval_kwh':[2.]*144, 'discharge_by_interval_kwh':[1.]*144,
                            'emergency_by_interval_kwh':[0.]*144,'soc_kwh':[6000.]*145})
        with tempfile.TemporaryDirectory() as tmp:
            path=export_scenario({'model':{'question':'4-3'},'daily':records},Path(tmp)/'out.xlsx')
            w=openpyxl.load_workbook(path,data_only=True)
            self.assertEqual(w['充放电量'].max_row,19)
            self.assertEqual(w['计划购电量']['B2'].value,0)
            self.assertEqual(w['计划购电量'].cell(2,145).value,143)
            self.assertEqual(w['充放电量'].cell(19,3).value,48)


if __name__=='__main__':unittest.main()
