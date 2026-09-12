from __future__ import annotations
import shutil,sys
from datetime import date,datetime
from pathlib import Path
import openpyxl,numpy as np
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT))
from src.CUMCM2026_C import paths,io_attachments as io
from src.CUMCM2026_C.solve_problem2_forecast import run
def main():
 data=run(); out=paths.result_path('result2'); shutil.copy2(paths.RESULT_TEMPLATES['result2'],out); wb=openpyxl.load_workbook(out); wp,wc,we=wb.worksheets; by={r['date']:r for r in data['daily']}
 for r,(d,rec) in enumerate(sorted(by.items()),2):
  dt=date.fromisoformat(d); wp.cell(r,1).value=datetime(dt.year,dt.month,dt.day)
  p=np.asarray(rec['purchase_by_interval_kwh']);
  for t in range(144): wp.cell(r,2+t).value=float(p[(t+1)%144])
  # Template's daily fee column is the normal planned-purchase fee; emergency
  # fees are reported separately in the emergency sheet/JSON summary.
  wp.cell(r,146).value=float(rec['plan_purchase_kwh']); wp.cell(r,147).value=float(rec['plan_cost_yuan'])
 # template charge sheet contains sample rows; fill rows when matching date exists
 for r in range(2,wc.max_row+1):
  v=wc.cell(r,1).value; d=v.date() if isinstance(v,datetime) else v
  key = d.isoformat() if isinstance(d, date) else str(d)[:10]
  if d is None or key not in by: continue
  rec=by[key]; b=(r-2)%6; ch=np.asarray(rec['charge_by_interval_kwh']); dis=np.asarray(rec['discharge_by_interval_kwh']); wc.cell(r,3).value=float(ch[24*b:24*(b+1)].sum()); wc.cell(r,4).value=float(dis[24*b:24*(b+1)].sum());
  if b==0: wc.cell(r,6).value=float(rec['soc_kwh'][0])
  if b==1: wc.cell(r,6).value=float(rec['soc_kwh'][-1])
 er=2
 for d,rec in sorted(by.items()):
  dt=date.fromisoformat(d)
  for t,q in enumerate(rec['emergency_by_interval_kwh'],1):
   if q>1e-9: we.cell(er,1).value=datetime(dt.year,dt.month,dt.day); we.cell(er,2).value=io.interval_label(t); we.cell(er,3).value=float(q); er+=1
 wb.save(out); print(out)
if __name__=='__main__':main()
