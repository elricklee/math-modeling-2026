from __future__ import annotations
import shutil, json, sys
from datetime import date, datetime
from pathlib import Path
import openpyxl, numpy as np
ROOT=Path(__file__).resolve().parents[1]; sys.path.insert(0,str(ROOT))
from src.CUMCM2026_C import paths, io_attachments as io
from src.CUMCM2026_C.solve_problem3 import forecast_vector, solve_rolling

def main():
 a1,a2,a3,a4=io.load_all(); price=np.asarray(a1.price.values,float); start=date.fromisoformat(paths.RESULT2_START)
 out=paths.result_path('result3'); out.parent.mkdir(parents=True,exist_ok=True); shutil.copy2(paths.RESULT_TEMPLATES['result3'],out)
 wb=openpyxl.load_workbook(out); ws_p=wb.worksheets[0]; ws_a=wb.worksheets[1]; ws_c=wb.worksheets[2]; ws_e=wb.worksheets[3]
 rows=[]
 for i,d in enumerate(a2.dates):
  if d<start: continue
  SCALE=0.825
  f=[(0,forecast_vector(a3,d,0,SCALE)),(36,forecast_vector(a3,d,6,SCALE)),(72,forecast_vector(a3,d,12,SCALE)),(108,forecast_vector(a3,d,18,SCALE))]
  ini,fin=solve_rolling(price,a2.load[i],f); p0=ini['purchase']; pf=fin['purchase']; pv=np.maximum(a2.pv_actual[i]*paths.INTERVAL_MINUTES/60-fin['curtail'],0); emg=np.maximum(a2.load[i]*paths.INTERVAL_MINUTES/60-(pf+fin['discharge']-fin['charge']+pv),0)
  rows.append((d,ini,fin,p0,pf,emg))
 for r,(d,ini,fin,p0,pf,emg) in enumerate(rows,2):
  ws_p.cell(r,1).value=datetime(d.year,d.month,d.day); ws_a.cell(r,1).value=datetime(d.year,d.month,d.day)
  for t in range(io.N_TIME_COLUMNS): ws_p.cell(r,2+t).value=float(p0[(t+1)%io.N_TIME_COLUMNS]); ws_a.cell(r,2+t).value=float(pf[(t+1)%io.N_TIME_COLUMNS])
  ws_p.cell(r,146).value=float(p0.sum()); ws_p.cell(r,147).value=float(np.dot(price,p0))
  ws_a.cell(r,146).value=float(pf.sum()); ws_a.cell(r,147).value=float(np.dot(price,pf))
 # charge rows: template has first 3 sample days; fill all matching dates if present
 by={d:(ini,fin) for d,ini,fin,_,_,_ in rows}
 for r in range(2,ws_c.max_row+1):
  v=ws_c.cell(r,1).value; d=v.date() if isinstance(v,datetime) else v
  if d not in by: continue
  ini,fin=by[d]; b=(r-2)%6; ws_c.cell(r,3).value=float(fin['charge'][24*b:24*(b+1)].sum()); ws_c.cell(r,4).value=float(fin['discharge'][24*b:24*(b+1)].sum());
  if b==0: ws_c.cell(r,6).value=float(fin['soc'][0])
  if b==1: ws_c.cell(r,6).value=float(fin['soc'][-1])
 er=2
 for d,ini,fin,p0,pf,emg in rows:
  for t,q in enumerate(emg,1):
   if q>1e-9:
    ws_e.cell(er,1).value=datetime(d.year,d.month,d.day); ws_e.cell(er,2).value=io.interval_label(t); ws_e.cell(er,3).value=float(q); er+=1
 wb.save(out); print(out)
if __name__=='__main__': main()
