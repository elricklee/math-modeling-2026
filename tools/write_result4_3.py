from __future__ import annotations
import shutil,sys
from datetime import date,datetime
from pathlib import Path
import openpyxl,numpy as np
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT))
from src.CUMCM2026_C import paths,io_attachments as io
from src.CUMCM2026_C.solve_problem3 import forecast_vector,solve_rolling
def main():
 a1,a2,a3,a4=io.load_all(); start=date.fromisoformat(paths.RESULT2_START); out=paths.result_path('result4-3'); shutil.copy2(paths.RESULT_TEMPLATES['result4-3'],out); wb=openpyxl.load_workbook(out); wp,wa,wc,we=wb.worksheets; rows=[]
 for i,d in enumerate(a2.dates):
  if d<start: continue
  p=np.asarray(a4.price[i],float); SCALE=0.825; f=[(0,forecast_vector(a3,d,0,SCALE)),(36,forecast_vector(a3,d,6,SCALE)),(72,forecast_vector(a3,d,12,SCALE)),(108,forecast_vector(a3,d,18,SCALE))]; ini,fin=solve_rolling(p,a2.load[i],f); pv=np.maximum(a2.pv_actual[i]/6-fin['curtail'],0); em=np.maximum(a2.load[i]/6-(fin['purchase']+fin['discharge']-fin['charge']+pv),0); rows.append((d,ini,fin,em,p))
 for r,(d,ini,fin,em,p) in enumerate(rows,2):
  for ws,v in ((wp,ini['purchase']),(wa,fin['purchase'])):
   ws.cell(r,1).value=datetime(d.year,d.month,d.day)
   for t in range(144): ws.cell(r,2+t).value=float(v[(t+1)%144])
   ws.cell(r,146).value=float(v.sum()); ws.cell(r,147).value=float(np.dot(p,v))
 er=2
 for d,ini,fin,em,p in rows:
  for t,q in enumerate(em,1):
   if q>1e-9: we.cell(er,1).value=datetime(d.year,d.month,d.day); we.cell(er,2).value=io.interval_label(t); we.cell(er,3).value=float(q); er+=1
 wb.save(out); print(out)
if __name__=='__main__': main()
