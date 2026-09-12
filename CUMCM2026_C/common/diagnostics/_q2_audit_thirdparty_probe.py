r"""t3 第三方复核：队长指定抽查项 + U1 量级裁定复核（独立计算）。

不复用编程手的任何代码；LP 与解析函数直接复用**我自己**前面写的
``_q2_audit_independent.py``（同属第三方复核，非被复核方代码）。

产出：``common/diagnostics/_q2_audit_thirdparty_probe.json``
"""

from __future__ import annotations

import json
import sys
from datetime import date
from pathlib import Path

import numpy as np
import openpyxl

DIAG = Path(__file__).resolve().parent
sys.path.insert(0, str(DIAG))

import _q2_audit_independent as A  # noqa: E402  我自己的复核模块

OUT = DIAG / "_q2_audit_thirdparty_probe.json"

# 队长在 t3 任务描述里给出的期望值
EXPECT_TOTAL_PURCHASE = 20218838.180253
EXPECT_TOTAL_COST = 12245046.915278
EXPECT_DAILY = {
    "2025-03-20": 66322.04493868313,
    "2025-06-21": 32667.359783333337,
    "2025-09-23": 66265.42130905349,
    "2025-12-21": 93770.21493971194,
}
EXPECT_Q1_COST = 35126.948589290
# 队长对 U1 的实测断言
EXPECT_U1 = {"att2_pv_peak_kw": 10216.20, "att2_pv_p999_kw": 9739.0,
             "att3_pv_peak_kw": 9995.89}


def read_attachment3_raw() -> np.ndarray:
    """附件 3 的全部预报数值（1460 x 24），前后向兼容地只取数值列。"""
    wb = openpyxl.load_workbook(A.RAW / "附件3.xlsx", data_only=True)
    try:
        ws = wb[wb.sheetnames[0]]
        rows = list(ws.iter_rows(values_only=True))
    finally:
        wb.close()
    vals = []
    for row in rows[1:]:
        nums = []
        for v in row[2:26]:
            try:
                nums.append(float(v))
            except (TypeError, ValueError):
                nums.append(float("nan"))
        if len(nums) == 24 and any(np.isfinite(nums)):
            vals.append(nums)
    return np.array(vals, dtype=float)


def main() -> int:
    """执行抽查与 U1 复核。"""
    results: dict[str, object] = {"spot_checks": [], "u1": {}, "pass": True}

    def rec(name: str, got: float, want: float, tol: float) -> None:
        ok = abs(got - want) <= tol
        results["pass"] = results["pass"] and bool(ok)
        results["spot_checks"].append(
            {"name": name, "got": float(got), "expected": float(want),
             "abs_gap": float(abs(got - want)), "tol": tol,
             "status": "passed" if ok else "failed"})
        print(f"[{'PASS' if ok else 'FAIL'}] {name}: got {got!r} vs expected {want!r} "
              f"(gap {abs(got - want):.3e}, tol {tol:g})")

    # ---------- 抽查：全年购电量与费用 ----------
    r2 = A.read_result2()
    plan_rows = r2["plan"]
    plan_vals = np.array([[A.num(v) for v in row[1:1 + A.T]] for row in plan_rows[1:]])
    tot_col = np.array([A.num(row[145]) for row in plan_rows[1:]])
    price1, load1, pv1 = A.read_attachment1()
    cost_col = np.array([A.num(row[146]) for row in plan_rows[1:]])

    rec("全年计划购电量（宽表 全天购电量 列求和）", float(tot_col.sum()),
        EXPECT_TOTAL_PURCHASE, 1e-3)
    rec("全年购电费（宽表 全天购电费 列求和）", float(cost_col.sum()),
        EXPECT_TOTAL_COST, 1e-3)
    rec("全年购电量（144 列全量求和，与上一行互为独立算式）",
        float(plan_vals.sum()), EXPECT_TOTAL_PURCHASE, 1e-3)

    # ---------- 抽查：四个指定日期全天购电量 ----------
    for iso, want in EXPECT_DAILY.items():
        d = date.fromisoformat(iso)
        i = A.TARGET_DATES.index(d)
        rec(f"{iso} 全天购电量", float(A.num(plan_rows[1 + i][145])), want, 1e-3)

    # ---------- 抽查：附件 1 特例复算（用我自己的 577 变量 LP） ----------
    r1 = A.solve_day(load1 / 6.0, pv1 / 6.0, price1)
    rec("附件 1 典型日复算费用（自建 LP）", float(r1["cost"]), EXPECT_Q1_COST, 1e-3)

    # ---------- 抽查：明细 CSV 与宽表交叉 ----------
    ts = A.read_timeseries()
    rec("明细 CSV 的 计划购电量 合计（与宽表互相独立）",
        float(ts["计划购电量(kWh)"].sum()), EXPECT_TOTAL_PURCHASE, 1e-3)

    # ---------- U1：附件 2 与附件 3 的光伏量级 ----------
    dates2, load2, pv2 = A.read_attachment2()
    att3 = read_attachment3_raw()
    daily_max = pv2.max(axis=1)
    u1 = {
        "att2_pv_peak_kw": float(pv2.max()),
        "att2_pv_p999_kw": float(np.quantile(pv2.ravel(), 0.999)),
        "att2_pv_daily_max_median_kw": float(np.median(daily_max)),
        "att2_pv_daily_max_min_kw": float(daily_max.min()),
        "att2_shape": list(pv2.shape),
        "att3_pv_peak_kw": float(np.nanmax(att3)),
        "att3_pv_p999_kw": float(np.nanquantile(att3, 0.999)),
        "att3_shape": list(att3.shape),
        "captain_claim": EXPECT_U1,
        "att2_peak_matches_captain": bool(abs(pv2.max() - EXPECT_U1["att2_pv_peak_kw"]) < 1.0),
        "att3_peak_matches_captain": bool(
            abs(float(np.nanmax(att3)) - EXPECT_U1["att3_pv_peak_kw"]) < 1.0),
    }
    results["u1"] = u1
    print("\n--- U1 量级复核 ---")
    for k, v in u1.items():
        print(f"  {k} = {v}")

    results["u1_closed"] = bool(u1["att2_peak_matches_captain"]
                                and u1["att3_peak_matches_captain"])
    print(f"\nU1 结论：{'关闭（附件 2 与附件 3 同尺度，编程手“实际约 1400 kW”不成立）' if results['u1_closed'] else '未能关闭，需人工判断'}")

    OUT.write_text(json.dumps(results, ensure_ascii=False, indent=2,
                              default=A._json_default), encoding="utf-8")
    print(f"\n写出：{OUT}")
    return 0 if results["pass"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
