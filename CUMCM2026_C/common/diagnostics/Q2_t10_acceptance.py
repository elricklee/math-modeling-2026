r"""t10 七项收口的验收自检（磁盘证据逐条核对，可被 captain / t8 复跑）。

定位
====

本脚本**不重算模型**，只核对 t10 七项收口在磁盘上的落地情况（脚本源码、两个 JSON、
result2.xlsx 的工作表形状）与"脚本 mtime 早于 JSON mtime""主口径逐位不变"这两条
一致性要求。逐条打印 ``PASS/FAIL``，全部 PASS 时退出码 0。

产出
====

* ``CUMCM2026_C/common/diagnostics/_q2_t10_acceptance.txt`` —— 本次核对的文本记录

运行方式（仓库根目录）::

    $env:PYTHONIOENCODING="utf-8"
    .venv\Scripts\python.exe CUMCM2026_C\common\diagnostics\Q2_t10_acceptance.py
"""

from __future__ import annotations

import json
import re
from datetime import datetime
from pathlib import Path

import openpyxl

_REPO_ROOT = Path(__file__).resolve().parents[3]
CASE = _REPO_ROOT / "CUMCM2026_C"
A_SRC = CASE / "Q2" / "Q2_analysis.py"
S_SRC = CASE / "Q2" / "Q2_solve_plan.py"
A_JSON = CASE / "Q2" / "outputs" / "Q2_analysis.json"
D_JSON = CASE / "Q2" / "outputs" / "Q2_diagnostics.json"
XLSX = CASE / "Q2" / "outputs" / "result2.xlsx"
AUD_SRC = CASE / "common" / "diagnostics" / "Q2_audit.py"
AUD_JSON = CASE / "common" / "diagnostics" / "_q2_audit.json"
OUT_TXT = CASE / "common" / "diagnostics" / "_q2_t10_acceptance.txt"

_LINES: list[str] = []
_OK_ALL = True


def chk(tag: str, cond: bool, detail: str) -> None:
    """记录一条核对结果（同时写屏幕与文本记录）。"""
    global _OK_ALL
    _OK_ALL = _OK_ALL and bool(cond)
    line = f"[{'PASS' if cond else 'FAIL'}] {tag}: {detail}"
    print(line)
    _LINES.append(line)


def mtime(p: Path) -> str:
    """文件的本地修改时间（用于"脚本早于产物"核对）。"""
    return datetime.fromtimestamp(p.stat().st_mtime).strftime("%Y-%m-%d %H:%M:%S")


def main() -> int:
    """核对七项收口，返回退出码（0 = 全部 PASS）。"""
    a_src = A_SRC.read_text(encoding="utf-8")
    s_src = S_SRC.read_text(encoding="utf-8")
    a_txt = A_JSON.read_text(encoding="utf-8")
    an = json.loads(a_txt)
    dg = json.loads(D_JSON.read_text(encoding="utf-8"))
    del s_src  # 仅用于 mtime 核对

    # --- 1) 双映射消除（唯一实质项）---------------------------------------
    # 只扫**赋值式**映射（避免把叙述文本里的契约公式误判为第二个实现）
    bad = re.findall(r"=\s*release_hour\s*\+\s*k(?!\s*-\s*1)", a_src)
    bad2 = re.findall(r"absolute\s*=\s*release_hour", a_src)
    chk("1 唯一映射实现", not bad and not bad2,
        f"赋值式 `= release_hour + k`（无 -1）出现 {len(bad)} 次、"
        f"`absolute = release_hour` 出现 {len(bad2)} 次")
    chk("1 唯一映射实现", a_src.count("def release_hourly_blocks") == 1
        and a_src.count("def forecast_from_release") == 1,
        f"release_hourly_blocks x{a_src.count('def release_hourly_blocks')}、"
        f"forecast_from_release x{a_src.count('def forecast_from_release')}")

    # --- 2) 同一发布时刻只有一个 MAE -------------------------------------
    mae_a = an["forecast_errors"]["pv_forecast_a_vs_actual"]["mae_kw"]
    br = an["forecast_errors"]["by_release_hour"]
    br0 = br["当日 00:00 发布"]["mae_kw"]
    chk("2 MAE 371.2208", abs(mae_a - 371.2208) < 1e-3, f"pv_forecast_a_vs_actual.mae_kw = {mae_a:.4f}")
    chk("2 by_release 不再矛盾", abs(br0 - mae_a) < 1e-9,
        f"by_release_hour['当日 00:00 发布'].mae_kw = {br0:.4f}（与主键相差 {abs(br0 - mae_a):.2e}）")
    chk("2 by_release 各口径齐全", sorted(br) == sorted([
        "当日 00:00 发布", "当日 06:00 发布", "当日 12:00 发布", "前一日 18:00 发布"]),
        f"键 {sorted(br)}")
    chk("2 0:00 行全天覆盖", br["当日 00:00 发布"]["target_day_covered_hours"] == 24
        and br["当日 00:00 发布"]["full_target_day_144_intervals"] is True
        and br["当日 00:00 发布"]["target_day_uncovered_hours"] == [],
        "0:00 发布：24/24 个小时块、144 时段全覆盖、未覆盖钟点为空")

    # --- 3) covered_hours_per_day 统一 + 陈旧文本清除 ---------------------
    cov = re.findall(r'"covered_hours_per_day":\s*([0-9]+)', a_src)
    chk("3 covered_hours_per_day", all(v == "24" for v in cov), f"脚本内取值 {cov}")
    chk("3 陈旧表述清除", "23 个小时块" not in a_src and "138/144" not in a_src
        and "t = 7..144" not in a_src and "mae_kw_covered_only" not in a_src
        and "covered_periods_only" not in a_src,
        "脚本内无『23 个小时块 / 138-144 / mae_kw_covered_only / covered_periods_only』")
    chk("3 JSON 无陈旧键", "mae_kw_covered_only" not in a_txt and "covered_periods_only" not in a_txt,
        "Q2_analysis.json 内无 mae_kw_covered_only / covered_periods_only")
    chk("3 JSON 覆盖定义", an["forecast_errors"]["covered_periods"]["full_day_intervals"] == 144
        and an["forecast_errors"]["covered_periods"]["covered_hours_per_day"] == 24,
        "covered_periods: full_day_intervals=144, covered_hours_per_day=24")

    # --- 4) forecast_definitions 措辞 ------------------------------------
    need = "覆盖当天 0:00–24:00 的全部 24 个小时块（144 个时段）"
    fa = an["forecast_definitions"]["预报A（主用）"]
    chk("4 预报A（主用）措辞", need in fa, f"{fa[:80]}...")

    # --- 5) 紧急购电量工作表仅表头（最终裁定）---------------------------
    wb = openpyxl.load_workbook(XLSX, data_only=False)
    ws3 = wb["紧急购电量"]
    shape3 = (ws3.max_row, ws3.max_column)
    comment = ws3["A1"].comment
    comment_txt = comment.text if comment is not None else None
    wb.close()
    chk("5 仅表头 max_row=1", shape3 == (1, 3), f"max_row={shape3[0]}, max_column={shape3[1]}")
    chk("5 A1 批注保留", bool(comment_txt) and "紧急购电量恒为 0" in (comment_txt or ""),
        f"批注长度 {len(comment_txt or '')}、含口径说明")
    chk("5 诊断 reason 保留", "reason" in dg["emergency_purchase"]
        and dg["emergency_purchase"]["main_convention_rows"] == 0,
        f"worksheet_shape={dg['emergency_purchase'].get('worksheet_shape_rows_cols')}, "
        f"header_only={dg['emergency_purchase'].get('worksheet_header_only')}")

    # --- 6) dispatch_bounds.note 按可复现性定性 --------------------------
    note = an["dispatch_bounds"]["note"]
    banned = [w for w in ("上界", "下界", "乐观", "保守") if w in note]
    chk("6 dispatch_bounds.note", not banned, f"违禁词 {banned or '无'}；note 前 60 字：{note[:60]}...")
    chk("6 可复现性定性", "复现" in note and "强假设" in note, "note 同时含『复现』与『强假设』")
    chk("6 JSON 无界式措辞", not any(w in a_txt for w in ("上界式", "乐观下界", "保守上界")),
        "Q2_analysis.json 内无『上界式/乐观下界/保守上界』")
    chk("6 承诺调度数值关系", an["dispatch_bounds"]["committed_emergency_kwh"]
        < an["dispatch_bounds"]["adaptive_emergency_kwh"],
        f"承诺调度 {an['dispatch_bounds']['committed_emergency_kwh']:,.2f} kWh "
        f"< 日内再调度 {an['dispatch_bounds']['adaptive_emergency_kwh']:,.2f} kWh"
        f"（故『承诺 = 上界』不成立，已按可复现性改写）")

    # --- 7) 顶层基线三字段 ----------------------------------------------
    exp = {"baseline_B0_purchase_kwh": 20997349.411166668,
           "baseline_B0_mean_price_yuan_per_kwh": 0.7813996,
           "purchase_reduction_vs_B0_kwh": 778511.2309135571}
    for k, v in exp.items():
        got = dg.get(k)
        chk("7 顶层基线字段", got is not None and abs(got - v) <= 1e-6 * max(1.0, abs(v)),
            f"{k} = {got!r}（期望 ≈ {v}，相对容差 1e-6）")
    chk("7 顶层与 metrics 同源", all(abs(dg[k] - dg["metrics"][k]) <= 1e-12 for k in exp),
        "顶层三字段与 metrics 内同名键逐位一致")

    # --- 8/9) 审计结论 ---------------------------------------------------
    aj = json.loads(AUD_JSON.read_text(encoding="utf-8"))
    u1 = next(u for u in aj["unverified_or_uncertain"] if u["id"] == "U1")
    chk("8 审计 20/20", aj["n_passed"] == 20 and aj["n_checks"] == 20
        and aj["verdict"] == "pass" and aj["n_failed"] == 0,
        f"{aj['n_passed']}/{aj['n_checks']} verdict={aj['verdict']}")
    chk("9 U1 已关闭", "已关闭" in u1["item"] and all(
        t in u1["problem"] for t in ("10,216.20", "9,739.22", "8,161.05", "9,995.8875", "1 400")),
        f"severity={u1['severity']}；含实测 10,216.20 / 9,739.22 / 8,161.05 / 9,995.8875 kW")

    # --- 10) mtime 与主口径不变 ------------------------------------------
    chk("10 mtime 脚本<JSON", A_SRC.stat().st_mtime < A_JSON.stat().st_mtime
        and S_SRC.stat().st_mtime < D_JSON.stat().st_mtime
        and AUD_SRC.stat().st_mtime < AUD_JSON.stat().st_mtime,
        f"analysis.py {mtime(A_SRC)} < analysis.json {mtime(A_JSON)}；"
        f"solve_plan.py {mtime(S_SRC)} < diagnostics.json {mtime(D_JSON)}；"
        f"audit.py {mtime(AUD_SRC)} < _q2_audit.json {mtime(AUD_JSON)}")
    m = dg["metrics"]
    chk("10 主口径逐位不变",
        abs(m["total_purchase_kwh"] - 20218838.180253) < 1e-6
        and abs(m["total_cost_yuan"] - 12245046.915278) < 1e-6
        and dg["sanity_checks"]["n_passed"] == 12
        and dg["sanity_checks"]["all_passed"] is True,
        f"购电 {m['total_purchase_kwh']:.6f} kWh、费用 {m['total_cost_yuan']:.6f} 元、"
        f"自检 {dg['sanity_checks']['n_passed']}/12")
    chk("10 主口径紧急购电 0", dg["emergency_purchase"]["main_convention_total_kwh"] == 0.0
        and dg["emergency_purchase"]["main_convention_rows"] == 0,
        "main_convention_total_kwh=0、rows=0")

    print()
    print("总判定：" + ("全部 PASS" if _OK_ALL else "存在 FAIL"))
    _LINES.append("")
    _LINES.append("总判定：" + ("全部 PASS" if _OK_ALL else "存在 FAIL"))
    OUT_TXT.parent.mkdir(parents=True, exist_ok=True)
    OUT_TXT.write_text("\n".join(_LINES) + "\n", encoding="utf-8")
    print(f"文本记录写出：{OUT_TXT}")
    return 0 if _OK_ALL else 1


if __name__ == "__main__":
    raise SystemExit(main())
