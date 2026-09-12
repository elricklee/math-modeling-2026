"""快速查看 C题 附件结构（表名、维度、前几行）。"""

from __future__ import annotations

import sys
from pathlib import Path

import openpyxl

ATTACH_DIR = Path(__file__).resolve().parents[1] / "CUMCM2026Problems" / "C题" / "附件"


def describe(path: Path, preview_rows: int = 6, preview_cols: int = 12) -> None:
    print("=" * 78)
    print(f"FILE: {path.name}  ({path.stat().st_size / 1024:.1f} KB)")
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    for ws in wb.worksheets:
        print("-" * 78)
        print(f"  SHEET: {ws.title!r}  dims={ws.dimensions}  "
              f"max_row={ws.max_row}  max_col={ws.max_column}")
        for i, row in enumerate(ws.iter_rows(max_row=preview_rows,
                                             max_col=preview_cols,
                                             values_only=True)):
            cells = ["" if c is None else str(c) for c in row]
            print(f"    r{i + 1}: {cells}")
    wb.close()


def main() -> None:
    targets = sorted(ATTACH_DIR.glob("*.xlsx")) + sorted((ATTACH_DIR / "附件5").glob("*.xlsx"))
    if len(sys.argv) > 1:
        wanted = set(sys.argv[1:])
        targets = [p for p in targets if p.name in wanted]
    for path in targets:
        describe(path)


if __name__ == "__main__":
    main()
