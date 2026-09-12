"""用附件4 的逐日电价反查附件1 对应的 2025 年日期（收敛阶段裁定用）。

方法：
    * 附件1 给出 144 个 10 分钟时段的电价；
    * 附件4 给出 365 天 × 144 时段的电价；
    * 对每一天计算与附件1 电价的逐时段平均绝对误差（MAE）与最大绝对误差，
      按 MAE 升序输出候选日期，用于判断附件1 是否来自附件4 的某一天。

附件1 的时间标签是"时段结束时刻"，所以它的第 t 个值与附件4 的同名时间点列对应；
本脚本同时检验两种对齐方式（按时间点等价、按区间偏移一格），以确认结论稳健。

用法：
    python tools/match_attachment1_date.py
"""

from __future__ import annotations

import json
from pathlib import Path

import numpy as np
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
RAW = ROOT / "data" / "raw" / "CUMCM2026_C"
OUT = ROOT / "data" / "processed" / "CUMCM2026_C"


def excel_serial_to_date(serial: float) -> str:
    """Excel 序列号（1900 日期系统）转 ISO 日期字符串。"""
    from datetime import datetime, timedelta

    base = datetime(1899, 12, 30)
    return (base + timedelta(days=float(serial))).strftime("%Y-%m-%d")


def load_attachment1() -> np.ndarray:
    """返回附件1 的 144 个电价。"""
    wb = openpyxl.load_workbook(RAW / "附件1.xlsx", read_only=True, data_only=True)
    ws = wb.active
    prices = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[1] is None:
            continue
        prices.append(float(row[1]))
    wb.close()
    return np.asarray(prices, dtype=float)


def load_attachment4() -> tuple[np.ndarray, list[str]]:
    """返回 (365×144 电价矩阵, 日期字符串列表)。"""
    wb = openpyxl.load_workbook(RAW / "附件4.xlsx", read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    header = rows[0]
    n_cols = len(header) - 1  # 去掉首列日期
    dates: list[str] = []
    matrix: list[list[float]] = []
    for row in rows[1:]:
        if row[0] is None:
            continue
        dates.append(excel_serial_to_date(row[0]) if isinstance(row[0], (int, float))
                     else str(row[0]))
        matrix.append([float(v) if v is not None else np.nan for v in row[1:1 + n_cols]])
    return np.asarray(matrix, dtype=float), dates


def main() -> None:
    a1 = load_attachment1()
    a4, dates = load_attachment4()
    print(f"附件1 电价点数: {len(a1)}   范围: {a1.min():.4f} – {a1.max():.4f}")
    print(f"附件4 矩阵: {a4.shape}   范围: {np.nanmin(a4):.4f} – {np.nanmax(a4):.4f}")
    print(f"附件4 日期: {dates[0]} … {dates[-1]}  共 {len(dates)} 天")

    results = []
    for idx, day in enumerate(a4):
        n = min(len(a1), len(day))
        # 对齐方式 A：逐时段直接比较
        mae_a = float(np.nanmean(np.abs(day[:n] - a1[:n])))
        max_a = float(np.nanmax(np.abs(day[:n] - a1[:n])))
        # 对齐方式 B：整日循环位移一格（检验时段标签口径差异）
        shifted = np.roll(day, 1)
        mae_b = float(np.nanmean(np.abs(shifted[:n] - a1[:n])))
        results.append((mae_a, max_a, mae_b, dates[idx]))

    results.sort(key=lambda item: item[0])
    print("\n按 MAE 升序的候选日期（前 8）：")
    print(f"{'日期':<12}{'MAE(直接)':>12}{'最大误差':>12}{'MAE(位移1格)':>16}")
    for mae_a, max_a, mae_b, day in results[:8]:
        print(f"{day:<12}{mae_a:>12.6f}{max_a:>12.6f}{mae_b:>16.6f}")

    exact = [r for r in results if r[1] < 1e-9]
    print(f"\n完全逐点相同的天数: {len(exact)}")
    for mae_a, max_a, _, day in exact[:5]:
        print(f"  精确匹配: {day}")

    best = results[0]
    payload = {
        "attachment1_points": int(len(a1)),
        "attachment1_price_min": float(a1.min()),
        "attachment1_price_max": float(a1.max()),
        "attachment4_shape": list(a4.shape),
        "attachment4_price_min": float(np.nanmin(a4)),
        "attachment4_price_max": float(np.nanmax(a4)),
        "attachment4_price_mean": float(np.nanmean(a4)),
        "best_match": {"date": best[3], "mae": best[0], "max_abs_error": best[1],
                       "mae_shifted_one_period": best[2]},
        "exact_match_dates": [r[3] for r in exact],
        "top8": [{"date": r[3], "mae": r[0], "max_abs_error": r[1],
                  "mae_shifted_one_period": r[2]} for r in results[:8]],
    }
    OUT.mkdir(parents=True, exist_ok=True)
    dest = OUT / "attachment1_date_match.json"
    dest.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"\n结果写入 {dest}")


if __name__ == "__main__":
    main()
